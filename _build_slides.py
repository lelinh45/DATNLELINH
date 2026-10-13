# -*- coding: utf-8 -*-
# Sinh slide bảo vệ DataIntegrity theo mẫu: 16:9, Arial, navy đậm + cyan accent
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY    = RGBColor(0x0E, 0x28, 0x41)
NAVY2   = RGBColor(0x15, 0x60, 0x82)
CYAN    = RGBColor(0x0F, 0x9E, 0xD5)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT   = RGBColor(0xC9, 0xD8, 0xE4)
ORANGE  = RGBColor(0xE9, 0x71, 0x32)
GREEN   = RGBColor(0x4E, 0xA7, 0x2E)
FONT    = 'Arial'

W, H = Emu(12192000), Emu(6858000)
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]

def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s

def bar(s, x, y, w, h, color=CYAN):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background()
    r.shadow.inherit = False
    return r

def txt(s, x, y, w, h, text, size=18, bold=False, color=WHITE,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.name = FONT; r.font.size = Pt(size)
        r.font.bold = bold; r.font.color.rgb = color
        r.font.italic = italic
    return tb

def bullets(s, x, y, w, h, items, size=16, gap=8):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (head, body) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = '▸ ' + head
        r.font.name = FONT; r.font.size = Pt(size + 1)
        r.font.bold = True; r.font.color.rgb = CYAN
        if body:
            p2 = tf.add_paragraph(); p2.space_after = Pt(gap)
            p2.line_spacing = 1.15
            r2 = p2.add_run(); r2.text = '   ' + body
            r2.font.name = FONT; r2.font.size = Pt(size - 1)
            r2.font.color.rgb = LIGHT
    return tb

def header(s, num, title):
    bar(s, Emu(457200), Emu(457200), Emu(274320), Emu(457200), ORANGE)
    txt(s, Emu(820420), Emu(400050), Emu(10200000), Emu(560000),
        (num + '. ' if num else '') + title, 26, True, WHITE)
    bar(s, Emu(457200), Emu(1036320), Emu(11277600), Emu(36576), NAVY2)

def footer(s, n):
    txt(s, Emu(10668000), Emu(6400000), Emu(1200000), Emu(300000),
        str(n).zfill(2), 12, False, LIGHT, PP_ALIGN.RIGHT)

def num_chip(s, x, y, n, color=NAVY2):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Emu(640080), Emu(460000))
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background(); r.shadow.inherit = False
    tf = r.text_frame
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    rr = p.add_run(); rr.text = str(n)
    rr.font.name = FONT; rr.font.size = Pt(16); rr.font.bold = True
    rr.font.color.rgb = WHITE

# ============ S1 — BÌA ============
s = slide()
bar(s, 0, 0, W, Emu(91440), CYAN)
txt(s, Emu(914400), Emu(700000), Emu(10363200), Emu(1400000),
    'BỘ CÔNG THƯƠNG\nTRƯỜNG ĐẠI HỌC ĐIỆN LỰC\nKHOA CÔNG NGHỆ THÔNG TIN',
    16, True, LIGHT, PP_ALIGN.CENTER)
bar(s, Emu(5029200), Emu(2350000), Emu(2133600), Emu(36576), ORANGE)
txt(s, Emu(914400), Emu(2650000), Emu(10363200), Emu(400000),
    'ĐỀ TÀI:', 20, True, ORANGE, PP_ALIGN.CENTER)
txt(s, Emu(914400), Emu(3100000), Emu(10363200), Emu(1900000),
    'TÌM HIỂU CÔNG NGHỆ BLOCKCHAIN VÀ ÁP DỤNG XÂY DỰNG HỆ THỐNG\nXÁC THỰC VÀ CHỐNG GIẢ MẠO DỮ LIỆU SỐ',
    26, True, WHITE, PP_ALIGN.CENTER)
txt(s, Emu(914400), Emu(5300000), Emu(10363200), Emu(400000),
    'DataIntegrity — Sổ cái phân tán · Ethereum Sepolia · IPFS · AES-256',
    14, False, CYAN, PP_ALIGN.CENTER)
bar(s, 0, Emu(6766560), W, Emu(91440), CYAN)

# ============ S2 — THỰC HIỆN ============
s = slide(); header(s, '', 'THỰC HIỆN ĐỀ TÀI'); footer(s, 2)
txt(s, Emu(1200000), Emu(1800000), Emu(9000000), Emu(500000),
    'Giảng viên hướng dẫn:', 18, True, CYAN)
txt(s, Emu(1200000), Emu(2300000), Emu(9000000), Emu(500000),
    'ThS. Đỗ Đức Cường', 22, True, WHITE)
bar(s, Emu(1200000), Emu(3100000), Emu(3200400), Emu(36576), NAVY2)
txt(s, Emu(1200000), Emu(3400000), Emu(9000000), Emu(500000),
    'Chủ nhiệm đề tài:', 18, True, CYAN)
txt(s, Emu(1200000), Emu(3900000), Emu(9000000), Emu(500000),
    'Lê Hồng Linh — 22810310058', 22, True, WHITE)
bar(s, Emu(1200000), Emu(4700000), Emu(3200400), Emu(36576), NAVY2)
txt(s, Emu(1200000), Emu(5000000), Emu(9000000), Emu(500000),
    'Thành viên tham gia (D17CNPM1):', 18, True, CYAN)
txt(s, Emu(1200000), Emu(5500000), Emu(10000000), Emu(1100000),
    'Nguyễn Xuân Hoàng Cường — 22810310036\nNguyễn Cao Nguyên — 22810310056\nNguyễn Quang Huy — 22810310051', 19, False, LIGHT)

# ============ S3 — TỔNG QUAN ============
s = slide(); header(s, '', 'TỔNG QUAN ĐỀ TÀI'); footer(s, 3)
agenda = [
    ('01', 'GIỚI THIỆU ĐỀ TÀI'),
    ('02', 'GIÁ TRỊ THỰC TIỄN'),
    ('03', 'MỤC TIÊU NGHIÊN CỨU'),
    ('04', 'KIẾN TRÚC & CÔNG NGHỆ'),
    ('05', 'KẾT QUẢ NGHIÊN CỨU'),
    ('06', 'Ý NGHĨA & HƯỚNG PHÁT TRIỂN'),
]
for i, (n, t) in enumerate(agenda):
    col, row = i % 2, i // 2
    x = Emu(1200000 + col * 5200000)
    y = Emu(1500000 + row * 1600000)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Emu(4600000), Emu(1200000))
    box.fill.solid(); box.fill.fore_color.rgb = NAVY2
    box.line.color.rgb = CYAN; box.line.width = Pt(1.25)
    box.shadow.inherit = False
    tf = box.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = n + '.'
    r.font.name = FONT; r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = ORANGE
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = t
    r2.font.name = FONT; r2.font.size = Pt(15); r2.font.bold = True; r2.font.color.rgb = WHITE

# ============ S4 — 01 GIỚI THIỆU ============
s = slide(); header(s, '01', 'GIỚI THIỆU ĐỀ TÀI'); footer(s, 4)
bullets(s, Emu(700000), Emu(1300000), Emu(10800000), Emu(5000000), [
    ('Bối cảnh thực tiễn',
     'Trong kỷ nguyên số hóa, hầu hết giá trị thông tin của tổ chức đều tồn tại dưới dạng dữ liệu số: văn bằng, chứng chỉ, hồ sơ bệnh án, hợp đồng, chứng từ kế toán, tài liệu nghiên cứu. Việc trao đổi và lưu trữ những dữ liệu này diễn ra hằng ngày, trên nhiều hệ thống khác nhau.'),
    ('Vấn đề đặt ra',
     'Khác với giấy tờ có dấu đỏ và con dấu thật, dữ liệu số có thể bị sao chép, chỉnh sửa hoặc giả mạo mà gần như không để lại dấu vết. Các cơ chế xác thực truyền thống đều phải tin cậy vào một bên thứ ba tập trung — điểm yếu chí mạng nếu bên đó bị tấn công hoặc cố tình sửa đổi.'),
    ('Hướng tiếp cận của đề tài',
     'Blockchain cung cấp một sổ cái phân tán, bất biến và minh bạch: một khi dấu vết (hash) của tài liệu được ghi lên chuỗi, không ai — kể cả quản trị hệ thống — có thể sửa hay xóa ngược nó. Đây chính là nền tảng để xây dựng cơ chế xác thực độc lập.'),
    ('Kết quả hướng tới',
     'Đề tài xây dựng hệ thống DataIntegrity: đóng dấu hash tài liệu lên mạng Ethereum, mã hóa nội dung bằng AES-256 và lưu trữ trên IPFS, cho phép bất kỳ ai có quyền cũng xác thực được tính toàn vẹn của tài liệu chỉ trong vài giây.'),
], 15, 10)

# ============ S5 — 02 GIÁ TRỊ ============
s = slide(); header(s, '02', 'GIÁ TRỊ THỰC TIỄN'); footer(s, 5)
bullets(s, Emu(700000), Emu(1300000), Emu(10800000), Emu(5000000), [
    ('Phát hiện giả mạo tức thì',
     'Thuật toán băm SHA-256 đảm bảo chỉ cần thay đổi một byte trong nội dung, hash của tài liệu sẽ khác hoàn toàn. Khi đối chiếu với hash đã đóng dấu trên chuỗi, mọi hành vi sửa đổi đều lộ rõ — không thể che giấu.'),
    ('Truy xuất nguồn gốc đầy đủ',
     'Mỗi bản ghi trên blockchain lưu vĩnh viễn ai là người tạo (địa chỉ ví), thời điểm ghi và hash của tài liệu. Hệ thống vì vậy trả lời được câu hỏi "tài liệu này do ai tạo, vào lúc nào" mà không cần tin bất kỳ ai.'),
    ('Chống chối bỏ trách nhiệm',
     'Mọi giao dịch ghi chuỗi đều được ký bằng khóa riêng của ví người dùng. Người tạo tài liệu không thể phủ nhận hành vi của mình — đây là nền tảng cho tính pháp lý của dữ liệu điện tử.'),
    ('Phạm vi ứng dụng',
     'Văn bằng và chứng chỉ học thuật, hồ sơ y tế, hợp đồng điện tử, hồ sơ hành chính công, chứng từ tài chính, hay chính các tài liệu nghiên cứu khoa học cần chứng minh thời điểm công bố.'),
], 15, 10)

# ============ S6 — 03 MỤC TIÊU ============
s = slide(); header(s, '03', 'MỤC TIÊU NGHIÊN CỨU'); footer(s, 6)
bullets(s, Emu(700000), Emu(1300000), Emu(10800000), Emu(5000000), [
    ('Nghiên cứu nền tảng lý thuyết',
     'Tìm hiểu kiến trúc Blockchain và mô hình đồng thuật, cơ chế hoạt động của Smart Contract trên nền tảng Ethereum, cùng các chuẩn mật mã học nền tảng: hàm băm SHA-256 và mã hóa đối xứng AES-256.'),
    ('Xây dựng hệ thống hoàn chỉnh',
     'Phát triển ứng dụng web cho phép người dùng tải tài liệu lên: hệ thống băm nội dung, mã hóa bằng khóa người dùng tự đặt, đưa ciphertext lên IPFS và ghi hash cùng metadata lên Smart Contract qua ví MetaMask.'),
    ('Đảm bảo an toàn và phân quyền',
     'Xác thực người dùng bằng chữ ký số thay mật khẩu, phê duyệt thiết bị đăng nhập mới, phân quyền nhiều cấp (Nhập liệu – Quản lý – Admin thực thi trong Smart Contract, thêm vai trò Người xem chỉ đọc ở tầng server).'),
    ('Vận hành và kiểm chứng thực tế',
     'Triển khai hệ thống đầy đủ trên Sepolia Testnet, MongoDB Atlas, IPFS Pinata và Render — đảm bảo hệ thống hoạt động 24/7 với bộ kiểm thử tự động bao phủ các luồng nghiệp vụ chính.'),
], 15, 10)

# ============ S7 — 04 KIẾN TRÚC ============
s = slide(); header(s, '04', 'KIẾN TRÚC HỆ THỐNG'); footer(s, 7)
def flow_box(s, x, y, w, h, title, sub, color=NAVY2):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.color.rgb = CYAN; r.line.width = Pt(1.25); r.shadow.inherit = False
    tf = r.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    rr = p.add_run(); rr.text = title
    rr.font.name = FONT; rr.font.size = Pt(14); rr.font.bold = True; rr.font.color.rgb = WHITE
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub
        r2.font.name = FONT; r2.font.size = Pt(10); r2.font.color.rgb = LIGHT
    return r

def arrow(s, x, y, w, h, vert=False):
    r = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW if vert else MSO_SHAPE.RIGHT_ARROW, x, y, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = CYAN
    r.line.fill.background(); r.shadow.inherit = False

flow_box(s, Emu(457200), Emu(2300000), Emu(3200000), Emu(1300000), 'TRÌNH DUYỆT NGƯỜI DÙNG', 'MetaMask · ethers.js · CryptoJS\nHash SHA-256 · Mã hóa AES-256')
arrow(s, Emu(3750000), Emu(2800000), Emu(500000), Emu(300000))
flow_box(s, Emu(4350000), Emu(2300000), Emu(3400000), Emu(1300000), 'BACKEND NODE.JS + EXPRESS', 'JWT · Phân quyền · Thông báo\nMongoDB Atlas')
arrow(s, Emu(7850000), Emu(2800000), Emu(500000), Emu(300000))
flow_box(s, Emu(8450000), Emu(1400000), Emu(3300000), Emu(1200000), 'ETHEREUM SEPOLIA', 'Smart Contract Solidity\nDocumentAdded · DocumentRevoked · RBAC', GREEN)
flow_box(s, Emu(8450000), Emu(3500000), Emu(3300000), Emu(1200000), 'IPFS / PINATA', 'Ciphertext đã mã hóa\nĐịnh danh bất biến bằng CID', NAVY2)
txt(s, Emu(8450000), Emu(2800000), Emu(3300000), Emu(400000), '▲▼', 16, True, CYAN, PP_ALIGN.CENTER)
txt(s, Emu(800000), Emu(4400000), Emu(10600000), Emu(2100000),
    '• Nội dung tài liệu chỉ rời trình duyệt sau khi ĐÃ được mã hóa — plaintext không bao giờ chạm internet; khóa mã hóa nằm ở người dùng hoặc được bọc bởi khóa chủ của server.\n'
    '• Blockchain đóng vai trò nguồn sự thật: hash, người tạo, thời điểm và trạng thái thu hồi. MongoDB chỉ lưu metadata phục vụ vận hành (tên file, dung lượng, chế độ chia sẻ).\n'
    '• Mọi yêu cầu kiểm chứng đều truy vấn trực tiếp Smart Contract — hệ thống không tin bất kỳ dữ liệu nào do phía client tự khai báo.',
    14, False, LIGHT)

# ============ S8 — CÔNG NGHỆ ============
s = slide(); header(s, '04', 'CÔNG NGHỆ SỬ DỤNG'); footer(s, 8)
bullets(s, Emu(700000), Emu(1350000), Emu(5300000), Emu(5000000), [
    ('Ethereum Sepolia Testnet', 'Mạng thử nghiệm Proof-of-Stake, mô phỏng đầy đủ mainnet với chi phí gas miễn phí — phù hợp triển khai và đánh giá thực tế.'),
    ('Smart Contract (Solidity)', 'Ghi nhận tài liệu theo lô, thu hồi kèm lý do, lưu trữ vai trò RBAC và phát sinh sự kiện truy vết.'),
    ('ethers.js v6 + MetaMask', 'Kết nối ví, ký giao dịch, đọc sự kiện chuỗi và kiểm tra quyền trực tiếp trên contract.'),
    ('AES-256 (CryptoJS)', 'Mã hóa nội dung file ngay trong trình duyệt trước khi đưa lên IPFS — dữ liệu rõ không bao giờ rời máy người dùng.'),
    ('SHA-256', 'Sinh dấu vết toàn vẹn của tài liệu; mọi thay đổi dù nhỏ nhất đều làm hash sai lệch hoàn toàn.'),
], 14, 8)
bullets(s, Emu(6300000), Emu(1350000), Emu(5200000), Emu(5000000), [
    ('IPFS / Pinata', 'Hệ thống lưu trữ phi tập trung; mỗi file được định danh bằng CID theo nội dung, không thể thay thế lén.'),
    ('Node.js 22 + Express', 'API xác thực, đồng bộ trạng thái, quản lý phiên thiết bị và thông báo điều hành.'),
    ('MongoDB Atlas', 'Lưu metadata tài liệu, hồ sơ ví và quyền, phiên thiết bị đăng nhập, API key của đối tác.'),
    ('JWT + chữ ký ví', 'Đăng nhập không mật khẩu: server phát thử thách, ví ký bằng khóa riêng, hệ thống cấp phiên JWT.'),
    ('Render + GitHub', 'Triển khai tự động mỗi khi push code — phiên bản mới luôn sẵn sàng phục vụ trên môi trường thật.'),
], 14, 8)

# ============ S9 — LUỒNG UPLOAD ============
s = slide(); header(s, '04', 'LUỒNG ĐÓNG DẤU TÀI LIỆU'); footer(s, 9)
steps = [
    'Chọn file\n(kéo-thả,\n≤100/lô)', 'Quét trùng\nsổ cái on-chain', 'Khóa AES +\nchế độ\ntừng file',
    'Băm SHA-256\n+ mã hóa file', 'Đẩy ciphertext\nlên IPFS → CID', 'MetaMask ký\naddDocuments', 'Ghi metadata\n→ MongoDB',
]
y = Emu(2050000)
for i, t in enumerate(steps):
    x = Emu(400000 + i * 1620000)
    flow_box(s, x, y, Emu(1400000), Emu(1100000), t, '', ORANGE if i in (1, 5) else NAVY2)
    if i < 6:
        arrow(s, Emu(1810000 + i * 1620000), Emu(2450000), Emu(200000), Emu(300000))
txt(s, Emu(700000), Emu(3700000), Emu(10800000), Emu(2700000),
    'Những điểm mấu chốt của thiết kế:\n'
    '• Trước khi nhập khóa, hệ thống đối chiếu hash từng file với sổ cái — bảng phân loại hiển thị ví và thời điểm đã đóng dấu: tài liệu đã thu hồi cho phép tải lại, tài liệu đang hoạt động chỉ tạo thêm bản ghi mới, không bao giờ ghi đè.\n'
    '• Popup bảo mật liệt kê rõ từng tệp được duyệt: một khóa AES chung cho cả lô, và mỗi file có thể đổi chế độ Công khai/Riêng tư hoặc đặt khóa riêng.\n'
    '• Một giao dịch duy nhất có thể đóng dấu tới 100 tài liệu (hàm addDocuments) — người dùng chỉ cần ký MetaMask một lần cho cả lô, tiết kiệm gas và thời gian thao tác.\n'
    '• Giao dịch blockchain được ghi xong mới lưu metadata vào MongoDB; nếu bước lưu bị lỗi, hàng đợi trong trình duyệt sẽ tự động thử lại ở phiên đăng nhập sau, tránh mất dấu tài liệu.',
    14, False, LIGHT)

# ============ S10 — LUỒNG XÁC THỰC ============
s = slide(); header(s, '04', 'XÁC THỰC & THU HỒI TÀI LIỆU'); footer(s, 10)
bullets(s, Emu(700000), Emu(1350000), Emu(5300000), Emu(4900000), [
    ('Cơ chế kiểm chứng', 'Người kiểm tra chọn file cần đối chiếu và nhập mã tài liệu. Hệ thống đọc hash gốc trực tiếp từ Smart Contract, đồng thời băm lại file hiện có và so sánh hai giá trị.'),
    ('Ba kết quả rõ ràng', 'VERIFIED — hash khớp, tài liệu nguyên bản. FAILED — hash sai lệch, tài liệu đã bị sửa đổi. REVOKED — nội dung nguyên bản nhưng đã bị vô hiệu hóa, không còn giá trị sử dụng.'),
    ('Truy xuất nguồn gốc', 'Mọi kết quả đều kèm thông tin ví người tạo và thời điểm ghi lên chuỗi, giúp truy vết đầy đủ vòng đời của tài liệu.'),
], 14, 8)
bullets(s, Emu(6400000), Emu(1350000), Emu(5100000), Emu(4900000), [
    ('Quyền thu hồi', 'Chỉ Admin và Quản lý được thu hồi; tài liệu do ví Admin tải lên thì chỉ Admin được thu hồi — quy tắc này được thực thi ngay trong Smart Contract, không thể lách từ giao diện.'),
    ('Ghi on-chain trước, đồng bộ sau', 'Sự kiện DocumentRevoked kèm lý do được ghi vĩnh viễn trên chuỗi trước, sau đó hệ thống mới cập nhật cờ trong MongoDB — blockchain luôn là nguồn chuẩn.'),
    ('Hoàn trả dung lượng', 'Khi tài liệu bị thu hồi, dung lượng lưu trữ của ví tải lên được hoàn trả tự động, phục vụ quản lý quota.'),
], 14, 8)

# ============ S11 — CHỨC NĂNG ============
s = slide(); header(s, '04', 'CÁC CHỨC NĂNG NỔI BẬT'); footer(s, 11)
bullets(s, Emu(700000), Emu(1350000), Emu(5300000), Emu(5000000), [
    ('Phân quyền RBAC trên chuỗi', 'Ba cấp quyền kế thừa lũy tiến NHẬP LIỆU (rank 1) → QUẢN LÝ (rank 2) → ADMIN (rank 3) được lưu và kiểm tra trực tiếp trên contract; vai trò NGƯỜI XEM chỉ đọc bổ sung ở tầng backend. Mọi quyền ghi do chính contract chặn — sửa database cũng không leo thang được quyền.'),
    ('Đăng nhập bằng chữ ký ví', 'Không dùng mật khẩu: server gửi một thử thách ngẫu nhiên, ví ký bằng khóa riêng, hệ thống xác minh và cấp phiên JWT — loại bỏ hoàn toàn nguy cơ rò rỉ mật khẩu.'),
    ('Phê duyệt thiết bị lạ', 'Thiết bị mới phải được một thiết bị đang hoạt động duyệt, hoặc OTP về email khi không còn máy nào online — mỗi phiên ghi nhãn trình duyệt kèm phiên bản, IP và User-Agent để chủ ví nhận diện máy lạ. Khi mất hết neo tin cậy, cơ chế chờ 24 giờ giữ cho ví không bị khóa vĩnh viễn.'),
], 14, 8)
bullets(s, Emu(6400000), Emu(1350000), Emu(5100000), Emu(5000000), [
    ('Chia sẻ Công khai / Riêng tư từng file', 'Chế độ hiển thị và khóa AES đặt riêng cho từng file trong cùng một lô. Tài liệu công khai: mọi ví đã đăng nhập được server cấp lại khóa và xem trực tiếp. Tài liệu riêng tư: chỉ chủ sở hữu và Admin được cấp khóa, người khác chỉ thấy metadata.'),
    ('Kiểm soát trùng nội dung hai lớp', 'Bảng phân loại trước khi nhập khóa: mục "đã thu hồi" cho tải lại, mục "đang hoạt động" cho ghi thêm bản ghi — mỗi dòng hiện ví và thời điểm đóng dấu on-chain. Lớp kiểm tra lần cuối còn bắt trường hợp nội dung vừa được đóng dấu sau khi quét.'),
    ('API key cho đối tác', 'Hệ thống bên ngoài có thể tải tài liệu tự động qua API key riêng; server tự sinh khóa AES-256, bọc bằng khóa chủ và ghi chuỗi như người dùng thường.'),
], 14, 8)

# ============ S12 — 05 KẾT QUẢ ============
s = slide(); header(s, '05', 'KẾT QUẢ NGHIÊN CỨU'); footer(s, 12)
stats = [
    ('2', 'Smart Contract\nđã triển khai trên Sepolia'),
    ('117', 'Test tự động\n79 unit + 38 e2e đều đạt'),
    ('24/7', 'Vận hành thực tế\nRender · Atlas · Pinata'),
    ('3+1', 'Vai trò RBAC\n3 trong contract + 1 chỉ-đọc ở server'),
]
for i, (n, t) in enumerate(stats):
    x = Emu(800000 + i * 2700000)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Emu(1500000), Emu(2400000), Emu(1500000))
    box.fill.solid(); box.fill.fore_color.rgb = NAVY2
    box.line.color.rgb = CYAN; box.line.width = Pt(1.25); box.shadow.inherit = False
    tf = box.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = n
    r.font.name = FONT; r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = ORANGE
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = t
    r2.font.name = FONT; r2.font.size = Pt(12); r2.font.color.rgb = LIGHT
bullets(s, Emu(700000), Emu(3400000), Emu(10800000), Emu(2900000), [
    ('Hệ thống chạy trên môi trường thật',
     'Contract 0xAf1e…dF74 đang hoạt động trên Sepolia; giao diện triển khai tại blockchain-data-integrity.onrender.com. Toàn bộ luồng tải lên — đóng dấu — kiểm chứng — thu hồi đều thao tác trực tiếp được.'),
    ('Phát hiện giả mạo ở mức bit',
     'Mọi chỉnh sửa dù nhỏ nhất đều bị phát hiện: verify so sánh SHA-256 của file nộp với bản ghi on-chain — đổi một ký tự là đổi toàn bộ dấu vân tay. Tài liệu đã thu hồi được cảnh báo riêng kèm lý do thay vì chỉ báo đúng/sai.'),
    ('Kiến trúc chống mất dữ liệu',
     'Metadata lưu MongoDB qua endpoint ghi lũy đẳng kèm hàng đợi tự phục hồi phía client; cache sự kiện on-chain chỉ tiến con trỏ trên đoạn quét liền kề nên bản ghi không bao giờ bị bỏ sót.'),
    ('Bảo mật nhiều lớp phối hợp',
     'Chữ ký ví thay mật khẩu, phiên JWT có thời hạn, duyệt thiết bị từ máy cũ kèm nhận diện IP–trình duyệt, OTP email làm đường cứu, phân quyền ngay trong contract, mã hóa AES-256 phía client và bọc khóa bằng khóa chủ phía server.'),
], 14, 8)

# ============ S13 — DEMO ============
s = slide(); header(s, '05', 'KỊCH BẢN TRÌNH DIỄN'); footer(s, 13)
steps = [
    ('1', 'Đăng nhập bằng MetaMask — ký thử thách, phê duyệt thiết bị lần đầu'),
    ('2', 'Kéo-thả nhiều tài liệu — bảng phân loại trùng trước, rồi đặt chế độ và khóa AES riêng cho từng file'),
    ('3', 'Theo dõi tiến trình: băm → mã hóa → IPFS → ký on-chain → đồng bộ MongoDB'),
    ('4', 'Xác thực file nguyên bản → VERIFIED; sửa 1 byte rồi verify lại → FAILED'),
    ('5', 'Thu hồi tài liệu — MetaMask ký giao dịch; verify lần nữa → REVOKED'),
    ('6', 'Phân quyền: khay chờ duyệt ví mới → cấp vai trò → đối chiếu trên sổ cái'),
]
for i, (n, t) in enumerate(steps):
    y = Emu(1350000 + i * 800000)
    num_chip(s, Emu(900000), y, n, ORANGE)
    txt(s, Emu(1700000), y + Emu(50000), Emu(9600000), Emu(700000), t, 15, False, WHITE)

# ============ S14 — 06 Ý NGHĨA ============
s = slide(); header(s, '06', 'Ý NGHĨA KHOA HỌC & HƯỚNG PHÁT TRIỂN'); footer(s, 14)
bullets(s, Emu(700000), Emu(1400000), Emu(5300000), Emu(4800000), [
    ('Ý nghĩa khoa học',
     'Đề tài chứng minh trên thực tế rằng blockchain hoàn toàn có thể đáp ứng bài toán chống giả mạo dữ liệu: toàn bộ vòng đời tài liệu — đóng dấu, kiểm chứng, thu hồi — đều được ghi nhận và kiểm tra công khai trên chuỗi.'),
    ('Đóng góp về mô hình',
     'Kiến trúc "hash và quyền on-chain, nội dung mã hóa off-chain" tách biệt rõ phần cần công khai và phần cần bảo mật, kèm cơ chế ủy thác khóa theo chế độ chia sẻ linh hoạt.'),
], 14, 8)
bullets(s, Emu(6400000), Emu(1400000), Emu(5100000), Emu(4800000), [
    ('Hạn chế còn tồn tại',
     'Hệ thống còn phụ thuộc độ ổn định của mạng Sepolia và chi phí gas; cơ chế ủy thác khóa tập trung tại server là điểm tin cậy đơn; hiệu năng với tệp dung lượng lớn chưa được tối ưu.'),
    ('Hướng phát triển tiếp theo',
     'Chuyển sang mạng Layer-2 để giảm chi phí ghi chuỗi; tích hợp chuẩn định danh phi tập trung DID và Verifiable Credential; áp dụng mã hóa ngưỡng (threshold) thay cho ủy thác khóa đơn điểm.'),
], 14, 8)

# ============ S15 — CẢM ƠN ============
s = slide()
bar(s, 0, 0, W, Emu(91440), CYAN)
txt(s, Emu(914400), Emu(2500000), Emu(10363200), Emu(1200000),
    'Cảm ơn thầy và các bạn đã lắng nghe', 36, True, WHITE, PP_ALIGN.CENTER)
bar(s, Emu(5029200), Emu(3900000), Emu(2133600), Emu(36576), ORANGE)
txt(s, Emu(914400), Emu(4200000), Emu(10363200), Emu(600000),
    'DataIntegrity — Xác thực và chống giả mạo dữ liệu số bằng Blockchain',
    16, False, CYAN, PP_ALIGN.CENTER)
bar(s, 0, Emu(6766560), W, Emu(91440), CYAN)

try:
    prs.save(r'D:\NCKH26\Thu-muc-lon\blockchain-data-integrity\LeHongLinh_Slides.pptx')
    print('Saved OK')
except PermissionError:
    prs.save(r'D:\NCKH26\Thu-muc-lon\blockchain-data-integrity\LeHongLinh_Slides_v2.pptx')
    print('File cu dang mo trong PowerPoint -> luu sang LeHongLinh_Slides_v2.pptx')
